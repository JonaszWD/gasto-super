"""One-off import of supermarket catalogue exports (.xlsx) into the price-comparison tables.

    uv run python -m collectors import-xlsx SupermaketData/products_dia_24092026-122757.xlsx
    uv run python -m collectors import-xlsx FILE... --dry-run        # parse only, no database

Expected sheet columns: Id, Nombre, Precio, Precio Pack, Formato, Categoria, Supermercado, Url, Url_imagen.
- Only Carrefour and Dia are imported; Mercadona rows are skipped (the live `mercadona` source
  covers it with EANs and per-warehouse prices).
- Prices are not tied to a postal code (stored with postal_code "") under source "xlsx".
- The observation time comes from the file name (DDMMYYYY-HHMMSS, Madrid time) or --observed-at,
  so re-importing an older export never overwrites newer prices.
- Ids are the chains' own product ids, the same ones EasyCompra uses: an EasyCompra listing for
  the same product lends its EAN (or its canonical product) so the two don't show up twice.
- Rows repeated under several categories are imported once; non-grocery categories are skipped.
"""

import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from sqlalchemy import Engine
from sqlmodel import Session, col, select

from app.models import CollectorRun, Listing, utcnow
from app.services.catalog import RawListing, merge_canonical, record_listing_price, upsert_listing, valid_ean
from app.services.money import to_cents
from app.services.quantity import Quantity, parse_quantity, unit_price_cents
from app.services.text import normalize
from app.sources.easycompra import guess_department
from collectors.runner import COMMIT_EVERY, RunStats

SOURCE = "xlsx"
IMPORT_CHAINS = {"Carrefour": "carrefour", "Dia": "dia"}
SKIPPED_CHAINS = {"Mercadona": "mercadona"}
COLUMNS = ("Id", "Nombre", "Precio", "Precio Pack", "Formato", "Categoria", "Supermercado", "Url", "Url_imagen")

# Formato (normalized) -> (base unit, factor turning the file's reference price into a per-base-unit price).
FORMATS: dict[str, tuple[str, float]] = {
    "kg": ("kg", 1.0),
    "kilo": ("kg", 1.0),
    "g": ("kg", 1000.0),
    "100g": ("kg", 10.0),
    "100 gr": ("kg", 10.0),
    "l": ("l", 1.0),
    "litro": ("l", 1.0),
    "100ml": ("l", 10.0),
    "100 ml": ("l", 10.0),
    "ud": ("unit", 1.0),
    "unidad": ("unit", 1.0),
    "lavado": ("unit", 1.0),
    "docena": ("unit", 1 / 12),
}

# Category path prefix -> department, or None to skip. Longest prefix wins. "food" means grocery:
# the name still decides food vs drink (as EasyCompra does) so the same product lands in the same
# department whichever source it came from.
CATEGORY_DEPARTMENTS: dict[str, dict[str, str | None]] = {
    "carrefour": {
        "/supermercado/bebe/alimentacion-infantil": "food",
        "/supermercado/bebe": None,
        "/supermercado/bebe-promocion": None,
        "/supermercado/bebidas": "drink",
        "/supermercado/congelados": "food",
        "/supermercado/congelados-promocion": "food",
        "/supermercado/cuidado-personal-e-higiene": None,
        "/supermercado/drogueria-y-limpieza": "household",
        "/supermercado/drogueria-y-limpieza/menaje": None,
        "/supermercado/drogueria-y-limpieza/papeleria": None,
        "/supermercado/el-mercado-promocion": "food",
        "/supermercado/frescos": "food",
        "/supermercado/la-despensa": "food",
        "/supermercado/la-despensa-promocion": "food",
        "/supermercado/limpieza-y-hogar-promocion": "household",
        "/supermercado/mascotas": None,
        "/supermercado/parafarmacia": None,
        "/supermercado/perfumeria-e-higiene-promocion": None,
    },
    "dia": {
        "/": "food",
        "/agua-y-refrescos": "drink",
        "/cervezas-vinos-y-licores": "drink",
        "/zumos-y-smoothies": "drink",
        "/limpieza-y-hogar": "household",
        "/cabello-y-perfumeria": None,
        "/higiene-y-cuidado-del-cuerpo": None,
        "/mascotas": None,
        "/salud-y-parafarmacia": None,
        "/infantil/higiene-y-cuidado": None,
        "/infantil/panales-y-toallitas": None,
        "/novedades-y-recomendados": "guess",
    },
}

_FILE_TS = re.compile(r"(\d{2})(\d{2})(\d{4})-(\d{2})(\d{2})(\d{2})")
MADRID = ZoneInfo("Europe/Madrid")


@dataclass
class ParseStats:
    rows: int = 0
    duplicates: int = 0
    skipped_category: int = 0
    skipped_chain: Counter[str] = field(default_factory=Counter)
    invalid: int = 0
    messages: list[str] = field(default_factory=list)


def observed_at_from_name(path: Path) -> datetime | None:
    m = _FILE_TS.search(path.name)
    if not m:
        return None
    d, mo, y, h, mi, s = (int(x) for x in m.groups())
    return datetime(y, mo, d, h, mi, s, tzinfo=MADRID)


def department_for(chain: str, category: str, name: str) -> str | None:
    rules = CATEGORY_DEPARTMENTS.get(chain, {})
    path = category.rstrip("/") + "/"
    best = max((p for p in rules if path.startswith(p.rstrip("/") + "/")), key=len, default=None)
    rule = rules[best] if best is not None else "food"
    if rule is None:
        return None
    guessed = guess_department(name)
    if rule == "guess":
        return guessed
    if rule == "food":
        return "drink" if guessed == "drink" else "food"
    return rule


def _is_promo(category: str) -> bool:
    return "-promocion/" in category


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cents(value: Any) -> int | None:
    try:
        return to_cents(value) if _text(value) else None
    except ValueError:
        return None


def resolve_quantity(name: str, price: int, reference: int | None, formato: str) -> tuple[Quantity | None, int | None]:
    """Pack size and unit price (cents per kg / l / unit).

    The name's size is preferred (exact); the chain's reference price is kept as the unit price when
    it agrees with it. Without a size in the name, the size is derived from price / reference price.
    """
    fmt = FORMATS.get(normalize(formato))
    file_unit_price = round(reference * fmt[1]) if fmt and reference and reference > 0 else None
    named = parse_quantity(name)
    if named is not None:
        computed = unit_price_cents(price, named)
        if fmt and named.unit == fmt[0] and file_unit_price and computed and 0.5 <= file_unit_price / computed <= 2:
            return named, file_unit_price
        return named, None  # different unit or implausible reference: computed from the size
    if fmt and file_unit_price:
        value = price / file_unit_price
        value = max(1.0, float(round(value))) if fmt[0] == "unit" else round(value, 3)
        if value > 0:
            return Quantity(value, fmt[0]), file_unit_price
    return None, None


def read_rows(path: Path) -> Iterator[dict[str, Any]]:
    from openpyxl import load_workbook  # dev dependency; only this CLI needs it

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = [_text(h) for h in next(rows, ())]
            missing = [c for c in COLUMNS if c not in header]
            if missing:
                raise ValueError(f"{path.name} / {ws.title}: missing columns {', '.join(missing)}")
            idx = {c: header.index(c) for c in COLUMNS}
            for row in rows:
                if row and any(v is not None for v in row):
                    yield {c: row[i] if i < len(row) else None for c, i in idx.items()}
    finally:
        wb.close()


def parse_files(
    paths: list[Path], chains: set[str] | None = None, observed_at: datetime | None = None
) -> tuple[list[RawListing], ParseStats]:
    stats = ParseStats()
    # One row per (chain, id); a product listed under a regular category wins over a promo one.
    chosen: dict[tuple[str, str], tuple[dict[str, Any], datetime]] = {}
    for path in paths:
        seen_at = observed_at or observed_at_from_name(path)
        if seen_at is None:
            raise ValueError(f"{path.name}: no DDMMYYYY-HHMMSS timestamp in the name; pass --observed-at")
        for row in read_rows(path):
            stats.rows += 1
            label = _text(row["Supermercado"])
            chain = IMPORT_CHAINS.get(label)
            if chain is None or (chains is not None and chain not in chains):
                stats.skipped_chain[chain or SKIPPED_CHAINS.get(label, label or "?")] += 1
                continue
            pid = _text(row["Id"])
            if not pid or not _text(row["Nombre"]) or not _cents(row["Precio"]):
                stats.invalid += 1
                if len(stats.messages) < 20:
                    stats.messages.append(f"{chain}/{pid or '?'}: missing id, name or price")
                continue
            key = (chain, pid)
            if key in chosen:
                stats.duplicates += 1
                if _is_promo(_text(chosen[key][0]["Categoria"])) and not _is_promo(_text(row["Categoria"])):
                    chosen[key] = (row, seen_at)
                continue
            chosen[key] = (row, seen_at)

    listings = []
    for (chain, pid), (row, seen_at) in chosen.items():
        name, category, price = _text(row["Nombre"]), _text(row["Categoria"]), _cents(row["Precio"])
        assert price is not None
        department = department_for(chain, category, name)
        if department is None:
            stats.skipped_category += 1
            continue
        quantity, unit_price = resolve_quantity(name, price, _cents(row["Precio Pack"]), _text(row["Formato"]))
        listings.append(
            RawListing(
                chain_id=chain,
                chain_product_id=pid[:80],
                name=name,
                price_cents=price,
                department=department,
                category=category.strip("/")[:120] or None,
                quantity=quantity,
                unit_price_cents=unit_price,
                image_url=_text(row["Url_imagen"])[:500] or None,
                url=_text(row["Url"])[:500] or None,
                observed_at=seen_at,
            )
        )
    return listings, stats


def _borrow_from_other_sources(session: Session, raw: RawListing) -> Listing | None:
    """The same chain product collected by another source (EasyCompra shares the ids)."""
    others = session.exec(
        select(Listing)
        .where(
            Listing.chain_id == raw.chain_id,
            Listing.chain_product_id == raw.chain_product_id,
            Listing.source != SOURCE,
        )
        .order_by(col(Listing.ean).is_(None), col(Listing.id))
    ).all()
    other = others[0] if others else None
    if other is not None and not raw.ean:
        raw.ean = valid_ean(other.ean)
    return other


def import_listings(engine: Engine, listings: list[RawListing], parse_stats: ParseStats) -> RunStats:
    chains = sorted({li.chain_id for li in listings})
    stats = RunStats(SOURCE, chains, "")
    with Session(engine) as session:
        run = CollectorRun(source=SOURCE, postal_code="")
        session.add(run)
        session.commit()
        for raw in listings:
            try:
                with session.begin_nested():
                    other = _borrow_from_other_sources(session, raw)
                    listing, is_new = upsert_listing(session, SOURCE, raw, "")
                    target = other.canonical_product_id if other else None
                    own = listing.canonical_product_id
                    if listing.link_source == "own" and target and own and own != target:
                        shared = session.exec(
                            select(Listing.id).where(Listing.canonical_product_id == own, Listing.id != listing.id)
                        ).first()
                        if shared is None:
                            merge_canonical(session, own, target)
                            session.refresh(listing)
                    changed = record_listing_price(session, listing, raw)
            except Exception as exc:  # one bad row must not stop the import
                stats.errors += 1
                if len(stats.messages) < 20:
                    stats.messages.append(f"{raw.chain_id}/{raw.chain_product_id}: {exc}")
                continue
            stats.products_checked += 1
            stats.new_listings += is_new
            stats.prices_changed += changed
            stats.eans_added += is_new and listing.ean is not None
            if stats.products_checked % COMMIT_EVERY == 0:
                session.commit()
        session.commit()

        stats.messages.extend(parse_stats.messages)
        stats.status = "partial" if stats.errors or parse_stats.invalid else "ok"
        if stats.products_checked == 0 and listings:
            stats.status = "failed"
        run.finished_at = utcnow()
        run.status = stats.status
        run.products_checked = stats.products_checked
        run.prices_changed = stats.prices_changed
        run.new_listings = stats.new_listings
        run.errors = stats.errors
        run.error_messages = "\n".join(stats.messages)[:4000]
        session.add(run)
        session.commit()
    return stats


def describe(listings: list[RawListing], stats: ParseStats) -> str:
    per_chain = Counter(li.chain_id for li in listings)
    per_dept = Counter(li.department for li in listings)
    no_qty = sum(li.quantity is None for li in listings)
    lines = [
        f"rows read: {stats.rows:,}",
        f"to import: {len(listings):,} ({', '.join(f'{c} {n:,}' for c, n in sorted(per_chain.items()))})",
        f"departments: {', '.join(f'{d} {n:,}' for d, n in sorted(per_dept.items()))}",
        f"without pack size: {no_qty:,}",
        f"duplicates merged: {stats.duplicates:,}",
        f"skipped (non-grocery category): {stats.skipped_category:,}",
        f"invalid rows: {stats.invalid:,}",
    ]
    for chain, n in sorted(stats.skipped_chain.items()):
        hint = " (use the live source)" if chain in SKIPPED_CHAINS.values() else ""
        lines.append(f"skipped chain {chain}: {n:,}{hint}")
    return "\n".join(lines)
